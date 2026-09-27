import json
import asyncio
import sqlite3
import zipfile
from pathlib import Path

from docx import Document

from app.services.files import extract_archive, read_submission, safe_member_path, student_name
from app.services.files import REVIEW_ITEM_MARKER, REVIEW_MARKER, write_reviewed_docx
from app.services.grading import calculate_score, extract_max_score, parse_response, summarize
from app import db
from app import main


def test_safe_archive_and_student_name(tmp_path: Path):
    archive = tmp_path / "homework.zip"
    with zipfile.ZipFile(archive, "w") as zipped:
        zipped.writestr("张三_001.txt", "答案")
    output = tmp_path / "files"
    files = extract_archive(archive, output)
    assert files[0].read_text(encoding="utf-8") == "答案"
    assert student_name(files[0]) == "张三 001"


def test_rejects_traversal():
    try:
        safe_member_path("../outside.txt")
    except ValueError:
        pass
    else:
        raise AssertionError("必须拒绝目录穿越路径")


def test_parse_and_summarize():
    result = parse_response(json.dumps({
        "score": 8,
        "max_score": 10,
        "deductions": [{"category": "概念遗漏", "reason": "缺少定义域"}],
    }))
    summary = summarize([result, result])
    assert result["score"] == 8
    assert summary["common_errors"][0] == {"category": "概念遗漏", "count": 2}


def test_parse_response_repairs_invalid_backslash_escapes():
    result = parse_response(r'{"score": 8, "max_score": 10, "evidence": "\theta"}')
    assert result["evidence"] == r"\theta"


def test_parse_response_keeps_latex_commands_with_valid_escape_prefixes():
    result = parse_response(r'{"evidence": "\theta \text{答案} \beta"}')
    assert result["evidence"] == r"\theta \text{答案} \beta"


def test_parse_response_keeps_valid_escapes():
    result = parse_response(r'{"message": "第一行\n第二行", "quote": "\"正确\""}')
    assert result["message"] == "第一行\n第二行"
    assert result["quote"] == '"正确"'


def test_calculate_score_from_deductions():
    result = calculate_score({
        "score": 55,
        "max_score": 100,
        "deductions": [{"points": 5}, {"points": 3}],
    })
    assert result["score"] == 92
    assert result["max_score"] == 100


def test_calculate_score_is_clamped():
    result = calculate_score({
        "score": 55,
        "max_score": 10,
        "deductions": [{"points": 12}],
    })
    assert result["score"] == 0


def test_extract_max_score_uses_rubric_or_defaults_to_100():
    assert extract_max_score("本次作业满分：90分") == 90
    assert extract_max_score("按评分细则批改") == 100


def test_docx_tables_are_read_and_annotated(tmp_path: Path):
    path = tmp_path / "table.docx"
    document = Document()
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "第1题：表格中的学生答案"
    document.save(path)

    content, _ = read_submission(path)
    assert "表格中的学生答案" in content
    write_reviewed_docx(
        path,
        8,
        10,
        [{"question": "第1题", "points": 2, "reason": "答案不完整", "evidence": "表格中的学生答案"}],
    )
    result = Document(path)
    assert any(paragraph.text.startswith(REVIEW_ITEM_MARKER) for paragraph in result.tables[0].cell(0, 0).paragraphs)


def test_settings_round_trip(tmp_path, monkeypatch):
    database = tmp_path / "settings.db"
    monkeypatch.setattr(db, "DATABASE_PATH", database)
    monkeypatch.setattr(db, "DATA_DIR", tmp_path)
    db.init_db()
    db.set_settings({"model": "test/model", "concurrency": "4"})
    assert db.get_settings(["model", "concurrency"]) == {
        "model": "test/model",
        "concurrency": "4",
    }


def test_init_db_migrates_rubric_source_path(tmp_path, monkeypatch):
    database = tmp_path / "homework.db"
    monkeypatch.setattr(db, "DATABASE_PATH", database)
    monkeypatch.setattr(db, "DATA_DIR", tmp_path)
    connection = sqlite3.connect(database)
    connection.execute(
        """
        CREATE TABLE rubrics (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL,
            description TEXT NOT NULL DEFAULT '',
            content TEXT NOT NULL,
            created_at TEXT NOT NULL,
            updated_at TEXT NOT NULL
        )
        """
    )
    connection.commit()
    connection.close()

    db.init_db()

    with db.connect() as connection:
        columns = {row["name"] for row in connection.execute("PRAGMA table_info(rubrics)")}
    assert "source_path" in columns


def test_delete_batch_removes_records_and_files(tmp_path, monkeypatch):
    database = tmp_path / "homework.db"
    upload_dir = tmp_path / "uploads"
    monkeypatch.setattr(db, "DATABASE_PATH", database)
    monkeypatch.setattr(db, "DATA_DIR", tmp_path)
    monkeypatch.setattr(main, "UPLOAD_DIR", upload_dir)
    main.running_tasks.clear()
    db.init_db()

    timestamp = db.now()
    with db.connect() as connection:
        rubric_id = connection.execute(
            "INSERT INTO rubrics (name, content, created_at, updated_at) VALUES (?, ?, ?, ?)",
            ("测试标准", "满分 10 分", timestamp, timestamp),
        ).lastrowid
        batch_id = connection.execute(
            """
            INSERT INTO batches
            (name, rubric_id, rubric_snapshot, status, created_at, updated_at)
            VALUES (?, ?, ?, 'completed', ?, ?)
            """,
            ("测试批次", rubric_id, "满分 10 分", timestamp, timestamp),
        ).lastrowid
        connection.execute(
            """
            INSERT INTO submissions
            (batch_id, filename, student_name, source_path, status, created_at, updated_at)
            VALUES (?, ?, ?, ?, 'completed', ?, ?)
            """,
            (batch_id, "作业.docx", "张三", str(upload_dir / str(batch_id) / "作业.docx"), timestamp, timestamp),
        )

    batch_dir = upload_dir / str(batch_id)
    batch_dir.mkdir(parents=True)
    (batch_dir / "作业.docx").write_text("test", encoding="utf-8")
    (batch_dir / "archive.zip").write_text("test", encoding="utf-8")

    assert asyncio.run(main.delete_batch(batch_id)) == {"status": "deleted"}
    assert not batch_dir.exists()
    with db.connect() as connection:
        assert connection.execute("SELECT 1 FROM batches WHERE id = ?", (batch_id,)).fetchone() is None
        assert connection.execute("SELECT 1 FROM submissions WHERE batch_id = ?", (batch_id,)).fetchone() is None


def test_write_reviewed_docx_replaces_previous_review(tmp_path: Path):
    path = tmp_path / "student.docx"
    document = Document()
    document.add_paragraph("第1题：学生原文答案")
    document.save(path)

    write_reviewed_docx(
        path,
        8,
        10,
        [{"question": "第1题", "points": 2, "reason": "概念遗漏", "evidence": "学生原文答案"}],
        ["请补充关键定义"],
    )
    write_reviewed_docx(
        path,
        9,
        10,
        [{"question": "第1题", "points": 1, "reason": "计算错误", "evidence": "学生原文答案"}],
    )

    result = Document(path)
    paragraphs = result.paragraphs
    assert paragraphs[0].text == "第1题：学生原文答案"
    assert paragraphs[1].text == f"{REVIEW_ITEM_MARKER} 第1题：扣分 1 分：计算错误"
    assert sum(item.text == REVIEW_MARKER for item in paragraphs) == 1
    assert "最终得分：9 / 10" in [item.text for item in paragraphs]
    assert "概念遗漏" not in "\n".join(item.text for item in paragraphs)
    review_paragraph = next(item for item in paragraphs if item.text.startswith(REVIEW_ITEM_MARKER))
    assert review_paragraph.runs[0].font.color.rgb is not None


def test_write_reviewed_docx_puts_unmatched_deduction_in_summary(tmp_path: Path):
    path = tmp_path / "student.docx"
    document = Document()
    document.add_paragraph("第1题：学生原文答案")
    document.save(path)

    write_reviewed_docx(
        path,
        7,
        10,
        [{"question": "第2题", "points": 3, "reason": "缺少步骤", "evidence": "不存在的原文"}],
    )

    text = "\n".join(paragraph.text for paragraph in Document(path).paragraphs)
    assert "未定位到原文的扣分项：" in text
    assert "第2题：扣分 3 分：缺少步骤" in text
