import base64
import io
import mimetypes
import re
import tempfile
import zipfile
from pathlib import Path

from ..config import SUPPORTED_SUFFIXES

REVIEW_MARKER = "【作业批改结果】"
REVIEW_ITEM_MARKER = "【批改】"


def _iter_document_paragraphs(document):
    from docx.oxml.table import CT_Tbl
    from docx.oxml.text.paragraph import CT_P
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    def walk(parent):
        element = parent._element.body if hasattr(parent._element, "body") else parent._tc
        for child in element.iterchildren():
            if isinstance(child, CT_P):
                yield Paragraph(child, parent)
            elif isinstance(child, CT_Tbl):
                table = Table(child, parent)
                for row in table.rows:
                    for cell in row.cells:
                        yield from walk(cell)

    yield from walk(document)


def safe_member_path(name: str) -> Path:
    path = Path(name)
    if path.is_absolute() or ".." in path.parts:
        raise ValueError("压缩包中包含不安全的路径")
    return path


def extract_archive(archive: Path, destination: Path) -> list[Path]:
    destination.mkdir(parents=True, exist_ok=True)
    files: list[Path] = []
    with zipfile.ZipFile(archive) as zipped:
        for member in zipped.infolist():
            if member.is_dir():
                continue
            relative = safe_member_path(member.filename)
            output = destination / relative
            output.parent.mkdir(parents=True, exist_ok=True)
            output.write_bytes(zipped.read(member))
            if output.suffix.lower() in SUPPORTED_SUFFIXES:
                files.append(output)
    return files


def student_name(path: Path) -> str:
    cleaned = re.sub(r"[_\-\s]+", " ", path.stem).strip()
    return cleaned or "未命名学生"


def read_submission(path: Path) -> tuple[str, str | None]:
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md"}:
        return path.read_text(encoding="utf-8", errors="replace"), None
    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        return "\n\n".join(page.extract_text() or "" for page in reader.pages), None
    if suffix == ".docx":
        from docx import Document

        document = Document(str(path))
        return "\n".join(paragraph.text for paragraph in _iter_document_paragraphs(document)), None
    if suffix in {".jpg", ".jpeg", ".png"}:
        mime = mimetypes.guess_type(path.name)[0] or "image/png"
        encoded = base64.b64encode(path.read_bytes()).decode("ascii")
        return "", f"data:{mime};base64,{encoded}"
    raise ValueError(f"暂不支持的文件类型：{suffix}")


def read_rubric(path: Path) -> str:
    if path.suffix.lower() in {".txt", ".md"}:
        return path.read_text(encoding="utf-8", errors="replace").strip()
    if path.suffix.lower() == ".docx":
        from docx import Document

        return "\n".join(
            paragraph.text for paragraph in _iter_document_paragraphs(Document(str(path)))
        ).strip()
    raise ValueError("评分规则只支持 DOCX、TXT 或 Markdown 文件")


def _remove_review_section(document) -> None:
    paragraphs = list(_iter_document_paragraphs(document))
    marker_index = next(
        (index for index, paragraph in enumerate(paragraphs) if paragraph.text.strip() == REVIEW_MARKER),
        None,
    )
    review_paragraphs = [paragraph for paragraph in paragraphs if paragraph.text.startswith(REVIEW_ITEM_MARKER)]
    if marker_index is not None:
        review_paragraphs.extend(paragraphs[marker_index:])
    for paragraph in review_paragraphs:
        paragraph._element.getparent().remove(paragraph._element)


def _normalized_text(value: str) -> str:
    return re.sub(r"\s+", "", value or "").casefold()


def _find_target_paragraph(paragraphs, deduction: dict):
    evidence = _normalized_text(str(deduction.get("evidence") or ""))
    question = _normalized_text(str(deduction.get("question") or ""))
    if evidence:
        for paragraph in paragraphs:
            if evidence in _normalized_text(paragraph.text):
                return paragraph
    if question:
        for paragraph in paragraphs:
            if question in _normalized_text(paragraph.text):
                return paragraph
    return None


def _insert_review_paragraph_after(paragraph, text: str, color):
    from docx.oxml import OxmlElement
    from docx.text.paragraph import Paragraph

    element = OxmlElement("w:p")
    paragraph._p.addnext(element)
    inserted = Paragraph(element, paragraph._parent)
    inserted.paragraph_format.left_indent = paragraph.paragraph_format.left_indent
    run = inserted.add_run(text)
    run.font.color.rgb = color
    return inserted


def write_reviewed_docx(
    path: Path,
    score: float | int | None,
    max_score: float | int | None,
    deductions: list[dict],
    comments: list[str] | None = None,
) -> None:
    from docx import Document
    from docx.shared import RGBColor

    document = Document(str(path))
    _remove_review_section(document)
    original_paragraphs = list(_iter_document_paragraphs(document))
    deductions_without_location = []
    insertion_points = {}

    for deduction in deductions:
        question = deduction.get("question") or "未标明位置"
        points = deduction.get("points", 0)
        reason = deduction.get("reason") or "未填写扣分原因"
        category = deduction.get("category")
        detail = f"{REVIEW_ITEM_MARKER} {question}：扣分 {points} 分：{reason}"
        if category:
            detail += f"（{category}）"
        target = _find_target_paragraph(original_paragraphs, deduction)
        if target is None:
            deductions_without_location.append(detail)
            continue
        anchor = insertion_points.get(id(target), target)
        insertion_points[id(target)] = _insert_review_paragraph_after(
            anchor,
            detail,
            RGBColor(192, 0, 0),
        )

    marker = document.add_paragraph()
    marker.paragraph_format.page_break_before = True
    marker.add_run(REVIEW_MARKER).bold = True
    score_paragraph = document.add_paragraph()
    score_run = score_paragraph.add_run(
        f"最终得分：{score if score is not None else '待定'}"
        f"{f' / {max_score}' if max_score is not None else ''}"
    )
    score_run.font.color.rgb = RGBColor(192, 0, 0)

    if deductions_without_location:
        unresolved = document.add_paragraph()
        unresolved_run = unresolved.add_run("未定位到原文的扣分项：")
        unresolved_run.font.color.rgb = RGBColor(192, 0, 0)
        for detail in deductions_without_location:
            paragraph = document.add_paragraph()
            run = paragraph.add_run(detail)
            run.font.color.rgb = RGBColor(192, 0, 0)

    for comment in comments or []:
        paragraph = document.add_paragraph()
        run = paragraph.add_run(f"教师评语：{comment}")
        run.font.color.rgb = RGBColor(192, 0, 0)

    with tempfile.NamedTemporaryFile(dir=path.parent, suffix=".docx", delete=False) as temporary:
        temporary_path = Path(temporary.name)
    try:
        document.save(str(temporary_path))
        temporary_path.replace(path)
    finally:
        temporary_path.unlink(missing_ok=True)


def archive_name(path: Path) -> str:
    return path.name
